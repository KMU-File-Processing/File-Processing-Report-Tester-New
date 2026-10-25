"""Result export: xlsx, deepthought-compatible text logs, JSON."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from .judge import PROJECTS, FINAL_BASE, FINAL_SCALE


def write_json(results, project: str, out: Path, extra: dict | None = None):
    payload = {
        "project": project,
        "generated_at": datetime.now().astimezone().isoformat(),
        "submissions": [r.as_dict() for r in results],
    }
    if extra:
        payload.update(extra)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                   encoding="utf-8")
    return out


def write_xlsx(results, project: str, out: Path):
    """Excel grade sheet. The final score is a formula so TAs can audit it."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError("openpyxl 이 필요합니다: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = PROJECTS.get(project, project)[:31]

    headers = ["학번", "이름", "제출 파일", "컴파일", "경고(-Werror)",
               "simple", "random", "메모리 누수", "원점수", "최종 점수"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    fill_bad = PatternFill("solid", fgColor="FFF0F0")
    fill_warn = PatternFill("solid", fgColor="FFF8E8")

    for i, r in enumerate(results, start=2):
        d = r.as_dict()
        sc = d["score"]
        comp = "성공" if d["compile"].get("ok") else "실패"
        werr = ("통과" if d["compile"].get("werror_ok") else "경고") \
            if d["compile"].get("ok") else "-"
        leak = {"yes": "발생", "no": "없음", "unknown": "확인 불가"}[d["leak"]]
        ws.append([d["student_id"], d["name"], d["source"], comp, werr,
                   sc["simple"], sc["random"], leak, sc["raw"], None])
        # final score as a formula
        ws.cell(row=i, column=10).value = f"={FINAL_BASE}+{FINAL_SCALE}*I{i}"
        if comp == "실패":
            for c in ws[i]:
                c.fill = fill_bad
        elif d["leak"] == "yes":
            for c in ws[i]:
                c.fill = fill_warn

    widths = [12, 12, 46, 10, 14, 10, 10, 12, 10, 12]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def write_deepthought(results, project: str, out_dir: Path):
    """Text logs compatible with the original grader's workflow.

    Writes one summary file and one detailed file per student.
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    def rule(title, width=80):
        t = f"= {title} "
        return t + "=" * max(0, width - len(t))

    summary = [rule(f"FP-Judge {PROJECTS.get(project, project)}"),
               f"생성 시각: {datetime.now().astimezone().isoformat()}",
               "",
               f"{'STUDENT_ID':<14}{'NAME':<12}{'COMPILE':<10}"
               f"{'SIMPLE':<10}{'RANDOM':<10}{'LEAK':<10}{'SCORE':<8}"]

    leak_label = {"yes": "Yes", "no": "No", "unknown": "?"}

    for r in results:
        d = r.as_dict()
        sc = d["score"]
        summary.append(
            f"{d['student_id']:<14}{d['name']:<12}"
            f"{('OK' if d['compile'].get('ok') else 'FAIL'):<10}"
            f"{sc['simple']:<10}{sc['random']:<10}"
            f"{leak_label[d['leak']]:<10}{sc['raw']:<8}")

        lines = [rule(f"{d['student_id']} {d['name']} / {project}"),
                 f"source: {d['source']}", ""]
        lines.append(rule("Compiling"))
        lines.append(f"std: {d['compile'].get('std', '-')}")
        lines.append(f"-Wall -Wextra -Werror: "
                     f"{'success' if d['compile'].get('werror_ok') else 'failure'}")
        if d["compile"].get("werror_stderr"):
            lines.append(d["compile"]["werror_stderr"])
        if d["compile"].get("stderr"):
            lines.append("")
            lines.append(d["compile"]["stderr"])
        lines.append("")

        for c in d["cases"]:
            lines.append(rule(f"TEST CASE - {c['name']}"))
            if c["verdict"] == "AC":
                lines.append("Diff OK :D")
            elif c["verdict"] == "WA" and c.get("diff"):
                lines.append(f"첫 불일치: {c['diff']['line']}행")
                for row in c["diff"]["rows"]:
                    if row["same"]:
                        lines.append(f"  {row['line']:>5}   {row['actual']}")
                    else:
                        lines.append(f"  {row['line']:>5} - {row['actual']}")
                        lines.append(f"  {row['line']:>5} + {row['expected']}")
                lines.append("")
                lines.append("Diff KO :(")
            else:
                lines.append(f"{c['verdict']}")
            lines.append("")

        lines.append(rule("RESULT"))
        lines.append(f"simple {sc['simple']}  random {sc['random']}  "
                     f"leak {d['leak']}  raw {sc['raw']}  final {sc['final']}")

        fname = f"{d['student_id']}_{d['name'] or 'unknown'}_deepthought.txt"
        (out_dir / fname).write_text("\n".join(lines), encoding="utf-8")

    (out_dir / "deepthought.txt").write_text("\n".join(summary), encoding="utf-8")
    return out_dir
